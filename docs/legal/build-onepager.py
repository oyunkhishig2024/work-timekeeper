import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT='/home/user/work-timekeeper/docs/legal/print/'
T={
'en':dict(
 file='Data-Protection-Summary-ONE-PAGE-Timekeeper-Work.docx',
 title='Timekeeper Work: how we protect personal data',
 sub='One-page summary for legal review · DRAFT · version dp-v1-draft · [date] · Operator: [name, registration no.]',
 what_h='What we hold and why',
 what='Attendance only: arrival and departure events at the workplace with the phone position at that moment, plus employee name, code, department and job title. No movement trail, no use other than attendance. Employees consent in writing; they see everything recorded about them in the app. Hosting: DigitalOcean, Singapore region, i.e. outside Mongolia (to be confirmed by counsel).',
 meas_h='Protection measures and their status',
 cols=('Area','Measure','Status'),
 st={'ok':'In place','prep':'Prepared, to be switched on','plan':'Planned before production','chk':'To be confirmed'},
 rows=[
  ('Minimisation','Only events leave the phone; exact coordinates are erased after 30 days; no trail is kept.','ok'),
  ('Retention','Attendance data deleted after 2 years, audit log after 24 months, by a daily automatic job. HR can irreversibly anonymise a former employee on request.','ok'),
  ('Encryption','HTTPS for all traffic. Backups encrypted with a key kept off the server.','ok'),
  ('Encryption at rest','Encryption of the live database disk at the hosting provider.','chk'),
  ('Separation','Each organisation is isolated inside the database (row-level security), so one cannot read another’s rows even after a code mistake.','ok'),
  ('Access control','Roles; managers see only their own branch or department; admins and HR use two-step sign-in; sessions expire; a phone account is tied to one registered phone.','ok'),
  ('Audit','Append-only log of who viewed or changed data; every export is logged.','ok'),
  ('Abuse protection','Rate limiting, account lockout, detection of impossible movement and duplicate devices. Cloudflare WAF in front of the server (needs the final domain).','ok|prep'),
  ('Server safety','Firewall on, 5-minute health alerts to the operator’s phone, written incident plan; the server was restarted and recovered by itself.','ok'),
  ('Backup and recovery','Nightly encrypted backups (last 14 kept), a daily whole-server copy at the hosting provider, and a restore tested on the server. Up to 24 h of data could be lost; point-in-time recovery to follow.','ok|plan'),
  ('Device integrity','Check that the app runs on a genuine, unmodified phone (Google Play Integrity).','plan'),
  ('Independent test','External security test and review before wide roll-out.','plan'),
 ],
 q_h='Questions for counsel',
 qs=['Is signed voluntary consent within an employment relationship a sufficient legal basis, or is a contract clause also needed?',
     'May the data be stored outside Mongolia (DigitalOcean, Singapore); what must the consent and the policy say?',
     'Are the retention periods (30 days coordinates, 2 years attendance, 24 months audit) acceptable?',
     'Breach duties: which authority, what deadline, who is informed? (Our incident plan has blanks for this.)',
     'Is registration or notification with the authority required? Which data-processing agreement do we need with each employer?'],
 foot='Sources in the project: docs/legal/privacy-policy-draft.*, docs/consent/consent-v2-draft.mn.txt, docs/operations/pilot-server-safety.md, docs/operations/monitoring-and-incidents.md, PRD 15.3.'),
'mn':dict(
 file='Мэдээлэл-хамгаалалт-ТӨСӨЛ-1-хуудас-Timekeeper-Work.docx',
 title='Timekeeper Work: хувийн мэдээллийг хэрхэн хамгаалах вэ',
 sub='Хуульчийн хяналтад зориулсан 1 хуудас · ТӨСӨЛ · хувилбар dp-v1-draft · [огноо] · Оператор: [нэр, регистрийн дугаар]',
 what_h='Бид юу хадгалж, яагаад',
 what='Зөвхөн ирц: ажлын байранд ирсэн, гарсан үйл явдал ба тэр мөчийн утасны байршил, мөн ажилтны нэр, код, хэлтэс, албан тушаал. Хөдөлгөөний зам хадгалахгүй, ирцээс өөр зорилгоор ашиглахгүй. Ажилтан бичгээр зөвшөөрч, өөрийнхөө тухай бүртгэсэн бүгдийг аппаас харна. Хостинг: DigitalOcean, Сингапур бүс, өөрөөр хэлбэл Монголын гадна (хуульч баталгаажуулна).',
 meas_h='Хамгаалах арга хэмжээ ба төлөв',
 cols=('Чиглэл','Арга хэмжээ','Төлөв'),
 st={'ok':'Хэрэгжсэн','prep':'Бэлэн, асаах хэрэгтэй','plan':'Production-оос өмнө','chk':'Баталгаажуулна'},
 rows=[
  ('Хамгийн бага мэдээлэл','Утаснаас зөвхөн үйл явдал гарна; яг координатыг 30 хоногийн дараа устгана; зам хадгалахгүй.','ok'),
  ('Хадгалах хугацаа','Ирцийн мэдээллийг 2 жилийн, аудитын бүртгэлийг 24 сарын дараа өдөр бүрийн автомат ажлаар устгана. Хүсэлтээр Хүний нөөц ажлаас гарсан ажилтныг эргэлт буцалтгүй нэргүйжүүлнэ.','ok'),
  ('Шифрлэлт','Бүх урсгал HTTPS. Нөөц хуулбарыг серверт байхгүй түлхүүрээр шифрлэнэ.','ok'),
  ('Хадгалалтын шифрлэлт','Үйлчилгээ үзүүлэгч дээрх ажиллаж буй өгөгдлийн сангийн дискийн шифрлэлт.','chk'),
  ('Тусгаарлалт','Байгууллага бүр өгөгдлийн санд тусгаарлагдсан (мөр түвшний хамгаалалт); кодын алдаа гарсан ч нэг нь нөгөөгийнхөө мөрийг уншихгүй.','ok'),
  ('Хандалтын хяналт','Эрхийн дүрэм; менежер зөвхөн өөрийн салбар/хэлтсийг харна; админ, Хүний нөөц хоёр шатлалт нэвтрэлттэй; сешн дуусдаг; утасны хаяг нэг бүртгэлтэй утсанд холбогддог.','ok'),
  ('Аудит','Хэн юу харсан, өөрчилсөн нэмэгдэх бүртгэл; экспорт бүр бүртгэгдэнэ.','ok'),
  ('Урвуулан ашиглахаас хамгаалах','Хурдны хязгаар, хаяг түгжих, боломжгүй хөдөлгөөн ба давхар төхөөрөмж илрүүлэх. Серверийн өмнө Cloudflare WAF (эцсийн domain хэрэгтэй).','ok|prep'),
  ('Серверийн аюулгүй байдал','Галт хана асаалттай, 5 минут тутмын эрүүл мэндийн мэдэгдэл операторын утсанд, бичмэл осол үйл явдлын төлөвлөгөө; сервер дахин ачаалагдахад өөрөө сэргэсэн.','ok'),
  ('Нөөц ба сэргээлт','Шөнө бүр шифрлэсэн нөөц (сүүлийн 14-ийг хадгална), үйлчилгээ үзүүлэгчийн өдөр бүрийн бүх серверийн хуулбар, сервер дээр шалгасан сэргээлт. 24 цаг хүртэлх мэдээлэл алдагдаж болно; цаг хугацааны цэгт сэргээлт дараа нь.','ok|plan'),
  ('Төхөөрөмжийн бүрэн бүтэн байдал','Апп жинхэнэ, өөрчлөөгүй утсан дээр ажиллаж байгааг шалгах (Google Play Integrity).','plan'),
  ('Бие даасан шалгалт','Өргөн нэвтрүүлэхээс өмнө гадны аюулгүй байдлын туршилт, хяналт.','plan'),
 ],
 q_h='Хуульчид тавих асуулт',
 qs=['Хөдөлмөрийн харилцаанд өгсөн гарын үсэгтэй сайн дурын зөвшөөрөл хууль зүйн хангалттай үндэслэл мөн үү, эсвэл гэрээний нэмэлт заалт хэрэгтэй юу?',
     'Мэдээллийг Монголын гадна (DigitalOcean, Сингапур) хадгалж болох уу; зөвшөөрөл, бодлогод юу бичих ёстой вэ?',
     'Хадгалах хугацаа (координат 30 хоног, ирц 2 жил, аудит 24 сар) зөвшөөрөгдөх үү?',
     'Зөрчлийн үүрэг: аль байгууллагад, хэдэн хоногт, хэнд мэдэгдэх вэ? (Манай төлөвлөгөөний энэ хэсэг хоосон.)',
     'Эрх бүхий байгууллагад бүртгүүлэх/мэдэгдэх шаардлага бий юу? Ажил олгогч бүртэй ямар мэдээлэл боловсруулах гэрээ байгуулах вэ?'],
 foot='Төслийн эх сурвалж: docs/legal/privacy-policy-draft.*, docs/consent/consent-v2-draft.mn.txt, docs/operations/pilot-server-safety.md, docs/operations/monitoring-and-incidents.md, PRD 15.3.'),
}
COLORS={'ok':'D9EAD3','prep':'FFF2CC','plan':'FCE5CD','chk':'FFF2CC'}
def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),fill); tcPr.append(sh)
def font(run,size,bold=False,color=None):
    run.font.name='Arial'; run._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial'); run._element.rPr.rFonts.set(qn('w:cs'),'Arial')
    run.font.size=Pt(size); run.bold=bold
    if color: run.font.color.rgb=RGBColor.from_string(color)
def para(d,text,size,bold=False,color=None,after=2,before=0):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(after); p.paragraph_format.space_before=Pt(before)
    r=p.add_run(text); font(r,size,bold,color); return p
for l,c in T.items():
    d=docx.Document()
    s=d.sections[0]; s.page_width=Cm(21); s.page_height=Cm(29.7)
    s.left_margin=s.right_margin=Cm(1.6); s.top_margin=Cm(1.3); s.bottom_margin=Cm(1.2)
    d.styles['Normal'].font.name='Arial'
    para(d,c['title'],17,True,after=1)
    para(d,c['sub'],8.5,color='555555',after=4)
    para(d,c['what_h'],11.5,True,after=1)
    para(d,c['what'],10,after=5)
    para(d,c['meas_h'],11.5,True,after=2)
    t=d.add_table(rows=1,cols=3); t.autofit=False
    widths=[Cm(3.3),Cm(11.0),Cm(3.5)]
    tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement(f'w:{e}'); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'999999'); b.append(x)
    tblPr.append(b)
    def setrow(row,vals,bold=False,fills=None):
        for i,(cell,v) in enumerate(zip(row.cells,vals)):
            cell.width=widths[i]; cell.text=''
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
            r=p.add_run(v); font(r,9,bold or i==0)
            if fills and fills[i]: shade(cell,fills[i])
    for i,w in enumerate(widths): t.columns[i].width=w
    setrow(t.rows[0],c['cols'],True,['E7E6E6']*3)
    for area,measure,status in c['rows']:
        keys=status.split('|')
        label=' / '.join(c['st'][k] for k in keys)
        row=t.add_row(); setrow(row,(area,measure,label),False,[None,None,COLORS[keys[-1] if len(keys)==1 else keys[-1]]])
    para(d,c['q_h'],11.5,True,after=1,before=5)
    for i,q in enumerate(c['qs'],1):
        p=para(d,f'{i}. {q}',10,after=2)
        p.paragraph_format.left_indent=Cm(0.4); p.paragraph_format.first_line_indent=Cm(-0.4)
    para(d,c['foot'],7.5,color='666666',before=4,after=0)
    d.save(OUT+c['file'])
